from docx import Document
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# === CREAR DOCUMENTO ===
doc = Document()

# === PORTADA ===
titulo = doc.add_heading("Módulo Profesional: Sistemas Informáticos", level=1)
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
info = [
    "Código: 0483",
    "Duración: 205 horas (10 ECTS)",
    "Profesor: Iván Jiménez Utiel",
    "Centro: IES Isidra de Guzmán",
    "Curso académico: 2025–2026"
]
for i in info:
    p = doc.add_paragraph(i)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

# === TABLA PRINCIPAL ===
doc.add_heading("Tabla resumen ministerial", level=2)
tabla = doc.add_table(rows=1, cols=4)
tabla.style = "Table Grid"
hdr = ["Unidad de Trabajo (UT)", "Contenidos Asociados", "RA Asociado", "Criterios de Evaluación (CE)"]
for i, h in enumerate(hdr):
    run = tabla.rows[0].cells[i].paragraphs[0].add_run(h)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 70, 140)

ut_data = [
    ["UT1. Evaluación de sistemas informáticos", 
     "Arquitectura de ordenadores. Von Neumann y Harvard. Componentes, periféricos, redes, topologías, normas de seguridad.", 
     "RA1", "RA1.a · RA1.b · RA1.c · RA1.d · RA1.e · RA1.f · RA1.g · RA1.h"],
    ["UT2. Instalación de sistemas operativos", 
     "Arquitectura y funciones del S.O. Tipos de sistemas, virtualización, instalación y actualización de S.O. y aplicaciones.", 
     "RA2", "RA2.a · RA2.b · RA2.c · RA2.d · RA2.e · RA2.f · RA2.g · RA2.h · RA2.i"],
    ["UT3. Gestión de la información y almacenamiento", 
     "Sistemas de archivos, particiones, RAID, copias de seguridad, tareas automáticas y administración de discos.", 
     "RA3", "RA3.a · RA3.b · RA3.c · RA3.d · RA3.e · RA3.f · RA3.g"],
    ["UT4. Administración de sistemas operativos", 
     "Usuarios y grupos, permisos, servicios, procesos, comandos y monitorización del sistema.", 
     "RA4", "RA4.a · RA4.b · RA4.c · RA4.d · RA4.e · RA4.f · RA4.g · RA4.h"],
    ["UT5. Interconexión de sistemas en red", 
     "Configuración de redes TCP/IP, IPv4/IPv6, redes cableadas e inalámbricas, protocolos y seguridad de red.", 
     "RA5", "RA5.a · RA5.b · RA5.c · RA5.d · RA5.e · RA5.f · RA5.g · RA5.h"],
    ["UT6. Gestión de sistemas en red y seguridad", 
     "Servidores de ficheros, impresión y aplicaciones. Conexiones remotas. Directivas de seguridad, cortafuegos y cifrado.", 
     "RA6", "RA6.a · RA6.b · RA6.c · RA6.d · RA6.e · RA6.f · RA6.g"],
    ["UT7. Documentación y software de propósito general", 
     "Software y licencias. Ofimática, correo, mensajería, trabajo colaborativo y documentación técnica.", 
     "RA7", "RA7.a · RA7.b · RA7.c · RA7.d · RA7.e · RA7.f · RA7.g"],
]
for ut in ut_data:
    row = tabla.add_row().cells
    for i, text in enumerate(ut):
        row[i].text = text

doc.add_paragraph()

# === PONDERACIÓN RAs ===
doc.add_heading("Ponderación de Resultados de Aprendizaje", level=2)
tabla2 = doc.add_table(rows=1, cols=3)
tabla2.style = "Table Grid"
hdr2 = ["RA", "Descripción resumida", "% del módulo"]
for i, h in enumerate(hdr2):
    r = tabla2.rows[0].cells[i].paragraphs[0].add_run(h)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 70, 140)

ponderaciones = [
    ("RA1", "Evalúa sistemas informáticos", "10 %"),
    ("RA2", "Instala sistemas operativos", "20 %"),
    ("RA3", "Gestiona la información del sistema", "15 %"),
    ("RA4", "Gestiona sistemas operativos", "15 %"),
    ("RA5", "Interconecta sistemas en red", "10 %"),
    ("RA6", "Opera sistemas en red", "20 %"),
    ("RA7", "Elabora documentación y usa software general", "10 %"),
]
for ra, desc, perc in ponderaciones:
    row = tabla2.add_row().cells
    row[0].text, row[1].text, row[2].text = ra, desc, perc

doc.add_page_break()

# === DESARROLLO UT ===
def ut_block(titulo, ra, ce, contenidos, metodologia, instrumentos, porcentaje):
    doc.add_heading(f"{titulo} ({porcentaje})", level=2)
    doc.add_paragraph(f"**Resultado de aprendizaje:** {ra}")
    doc.add_paragraph(f"**Criterios de evaluación:** {ce}")
    doc.add_paragraph(f"**Contenidos asociados:** {contenidos}")
    doc.add_paragraph(f"**Actividades y metodología:** {metodologia}")
    doc.add_paragraph(f"**Instrumentos de evaluación:** {instrumentos}")
    doc.add_paragraph("")

ut_block("UT1. Evaluación de sistemas informáticos",
         "Evalúa sistemas informáticos identificando sus componentes y características.",
         "RA1.a a RA1.h",
         "Arquitectura, componentes, periféricos, redes y seguridad.",
         "Tareas prácticas de identificación de hardware y simulación de redes locales.",
         "Pruebas escritas, prácticas de laboratorio y observación directa.", "10 %")

ut_block("UT2. Instalación de sistemas operativos",
         "Instala sistemas operativos planificando el proceso e interpretando documentación técnica.",
         "RA2.a a RA2.i",
         "Instalación y configuración de S.O. libres y propietarios, virtualización, gestores de arranque, licencias.",
         "Instalaciones guiadas, uso de máquinas virtuales, prácticas de documentación técnica.",
         "Cuaderno de prácticas y prueba práctica de instalación completa.", "20 %")

ut_block("UT3. Gestión de la información y almacenamiento",
         "Gestiona la información del sistema aplicando medidas de seguridad e integridad.",
         "RA3.a a RA3.g",
         "Particiones, RAID, copias de seguridad, tareas automáticas y administración de discos.",
         "Trabajo práctico con discos virtuales, copias programadas y recuperación de datos.",
         "Ejercicios prácticos y prueba de laboratorio.", "15 %")

ut_block("UT4. Administración de sistemas operativos",
         "Gestiona sistemas operativos utilizando comandos y herramientas gráficas.",
         "RA4.a a RA4.h",
         "Cuentas, permisos, servicios, procesos y monitorización.",
         "Prácticas en entorno gráfico y consola.",
         "Rúbrica de administración y observación del trabajo en equipo.", "15 %")

ut_block("UT5. Interconexión de sistemas en red",
         "Interconecta sistemas configurando dispositivos y protocolos.",
         "RA5.a a RA5.h",
         "Configuración TCP/IP, IPv6, redes cableadas e inalámbricas, seguridad básica.",
         "Prácticas de configuración y pruebas con herramientas de diagnóstico.",
         "Pruebas de red y evaluación por desempeño.", "10 %")

ut_block("UT6. Gestión de sistemas en red y seguridad",
         "Opera sistemas en red gestionando sus recursos y aplicando seguridad.",
         "RA6.a a RA6.g",
         "Servidores, directivas, cifrado, cortafuegos y dominios.",
         "Escenarios de simulación de servidor y configuración de permisos.",
         "Pruebas prácticas y rúbricas de seguridad.", "20 %")

ut_block("UT7. Documentación y software de propósito general",
         "Elabora documentación utilizando aplicaciones informáticas de propósito general.",
         "RA7.a a RA7.g",
         "Ofimática, trabajo colaborativo, correo, mensajería y búsqueda de documentación técnica.",
         "Trabajos escritos, presentaciones y búsquedas documentales.",
         "Entregas de documentación y evaluación de presentaciones.", "10 %")

# === TABLA RA-CONTENIDOS-HORAS ===
doc.add_page_break()
doc.add_heading("Relación RA – Contenidos – Horas", level=2)
tabla3 = doc.add_table(rows=1, cols=4)
tabla3.style = "Table Grid"
hdr3 = ["RA", "Unidad/es de Trabajo asociadas", "Contenidos principales", "Horas estimadas"]
for i, h in enumerate(hdr3):
    r = tabla3.rows[0].cells[i].paragraphs[0].add_run(h)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 70, 140)

data3 = [
    ("RA1", "UT1", "Evaluación de hardware y redes locales", "20 h"),
    ("RA2", "UT2", "Instalación y configuración de S.O.", "40 h"),
    ("RA3", "UT3", "Gestión de la información y almacenamiento", "30 h"),
    ("RA4", "UT4", "Administración de S.O.", "30 h"),
    ("RA5", "UT5", "Interconexión y configuración de redes", "20 h"),
    ("RA6", "UT6", "Gestión de sistemas en red y seguridad", "40 h"),
    ("RA7", "UT7", "Documentación y software general", "25 h"),
]
for row_data in data3:
    row = tabla3.add_row().cells
    for i, text in enumerate(row_data):
        row[i].text = text

# === GUARDAR ===
doc.save("Sistemas_Informaticos_UTs_2025_2026.docx")
print("✅ Documento generado correctamente: Sistemas_Informaticos_UTs_2025_2026.docx")
